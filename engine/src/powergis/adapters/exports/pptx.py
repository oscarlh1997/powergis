"""Exportación a PowerPoint (fase 8).

Una portada, un resumen ejecutivo, una diapositiva por subsección con sus KPIs
y su tabla resumida, el PlaceRank y una de fuentes. Deliberadamente sobrio:
es material para llevar a una reunión, no una presentación de marketing.
"""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ...domain.models import Report

if TYPE_CHECKING:  # pragma: no cover
    # `pptx.Presentation` es una función fábrica, no una clase: usarla como
    # anotación es incorrecto. El tipo de verdad es este.
    from pptx.presentation import Presentation as Deck
else:
    Deck = Any

INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x1E, 0x29, 0x3B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0xF0, 0xFD, 0xF4)
RED = RGBColor(0xFE, 0xF2, 0xF2)

MAX_TABLE_ROWS = 12
MAX_TABLE_COLS = 7


class PptxExporter:
    extension = "pptx"
    media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def render(self, report: Report) -> bytes:
        deck = Presentation()
        deck.slide_width = Inches(13.333)
        deck.slide_height = Inches(7.5)

        self._cover(deck, report)
        if report.executive_summary:
            self._summary(deck, report)

        for section in report.sections:
            if not section.available:
                continue
            for sub in section.subsections:
                self._subsection(deck, section.title, sub)

        if report.placerank:
            self._placerank(deck, report.placerank)

        self._sources(deck, report)

        buffer = io.BytesIO()
        deck.save(buffer)
        return buffer.getvalue()

    # ------------------------------------------------------------------ #

    def _blank(self, deck: Deck):
        return deck.slides.add_slide(deck.slide_layouts[6])

    def _title(self, slide, text: str, subtitle: str | None = None) -> None:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
        frame = box.text_frame
        frame.text = text
        run = frame.paragraphs[0].runs[0]
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.color.rgb = INK
        if subtitle:
            para = frame.add_paragraph()
            para.text = subtitle
            para.runs[0].font.size = Pt(12)
            para.runs[0].font.color.rgb = MUTED

    def _cover(self, deck: Deck, report: Report) -> None:
        slide = self._blank(deck)
        box = slide.shapes.add_textbox(Inches(0.9), Inches(2.2), Inches(11.5), Inches(3))
        frame = box.text_frame
        frame.text = "Análisis de localización comercial"
        frame.paragraphs[0].runs[0].font.size = Pt(40)
        frame.paragraphs[0].runs[0].font.bold = True
        frame.paragraphs[0].runs[0].font.color.rgb = INK

        for label, value in [
            ("Ámbito", f"{report.scope.get('name', '')} ({report.scope.get('level', '')})"),
            ("Zonas comparadas", report.scope.get("children_count", "")),
            ("Sector", report.business.get("sector") or "—"),
            ("Informe", f"{report.tier} · v{report.version}"),
            ("Generado", report.generated_at.strftime("%d/%m/%Y")),
        ]:
            para = frame.add_paragraph()
            para.text = f"{label}:  {value}"
            para.runs[0].font.size = Pt(14)
            para.runs[0].font.color.rgb = MUTED

    def _summary(self, deck: Deck, report: Report) -> None:
        slide = self._blank(deck)
        self._title(slide, "Resumen ejecutivo")
        summary = report.executive_summary or {}
        self._kpi_row(slide, [
            {"label": h["label"], "value": h.get("value"), "unit": h.get("unit", "")}
            for h in summary.get("headline", [])
        ])

        best = summary.get("mejores_zonas") or []
        worst = summary.get("peores_zonas") or []
        box = slide.shapes.add_textbox(Inches(0.6), Inches(3.1), Inches(12.1), Inches(3.6))
        frame = box.text_frame
        frame.word_wrap = True

        narrative = summary.get("narrative")
        if isinstance(narrative, dict) and narrative.get("resumen"):
            frame.text = narrative["resumen"]
            frame.paragraphs[0].runs[0].font.size = Pt(13)
        else:
            frame.text = ""

        for title, zones in (("Mejores zonas", best), ("Zonas con menor puntuación", worst)):
            if not zones:
                continue
            para = frame.add_paragraph()
            para.text = title
            para.runs[0].font.bold = True
            para.runs[0].font.size = Pt(13)
            for zone in zones[:3]:
                item = frame.add_paragraph()
                score = zone.get("score")
                item.text = f"· {zone.get('zona', '')}" + (f" — {score}" if score else "")
                item.runs[0].font.size = Pt(12)
                item.runs[0].font.color.rgb = MUTED

    def _subsection(self, deck: Deck, section_title: str, sub) -> None:
        slide = self._blank(deck)
        self._title(slide, sub.title, section_title)
        self._kpi_row(slide, [
            {"label": k.label, "value": k.value, "unit": str(k.unit), "decimals": k.decimals}
            for k in sub.kpis[:5]
        ])
        if sub.tables:
            self._table(slide, sub.tables[0], top=Inches(3.0))

    def _kpi_row(self, slide, kpis: list[dict[str, Any]]) -> None:
        if not kpis:
            return
        width = Inches(2.3)
        gap = Inches(0.22)
        # `Inches()` devuelve Emu (un int); el acumulador se anota como
        # int porque abajo se le suman anchos y separaciones.
        left: int = Inches(0.6)
        for kpi in kpis[:5]:
            box = slide.shapes.add_textbox(left, Inches(1.55), width, Inches(1.15))
            frame = box.text_frame
            frame.word_wrap = True
            frame.text = str(kpi["label"])[:38]
            frame.paragraphs[0].runs[0].font.size = Pt(9)
            frame.paragraphs[0].runs[0].font.color.rgb = MUTED

            para = frame.add_paragraph()
            value = kpi.get("value")
            if value is None:
                para.text = "no disponible"
                para.runs[0].font.size = Pt(13)
                para.runs[0].font.italic = True
                para.runs[0].font.color.rgb = MUTED
            else:
                para.text = f"{_fmt(value, kpi.get('decimals', 2))} {kpi.get('unit', '')}"
                para.runs[0].font.size = Pt(19)
                para.runs[0].font.bold = True
                para.runs[0].font.color.rgb = INK
            left += width + gap

    def _table(self, slide, table, top) -> None:
        columns = table.columns[:MAX_TABLE_COLS]
        rows = table.rows[:MAX_TABLE_ROWS]
        if not rows:
            return

        shape = slide.shapes.add_table(
            len(rows) + 1, len(columns), Inches(0.6), top,
            Inches(12.1), Inches(0.32 * (len(rows) + 1)),
        )
        grid = shape.table
        tops = set(table.highlights.get("top", []))
        bottoms = set(table.highlights.get("bottom", []))

        for c, column in enumerate(columns):
            cell = grid.cell(0, c)
            cell.text = column.label[:26]
            para = cell.text_frame.paragraphs[0]
            para.runs[0].font.size = Pt(9)
            para.runs[0].font.bold = True
            para.runs[0].font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT

        for r, row in enumerate(rows, start=1):
            code = row.get("geo_code")
            highlight = GREEN if code in tops else (RED if code in bottoms else None)
            for c, column in enumerate(columns):
                cell = grid.cell(r, c)
                value = row.get(column.key)
                if value is None:
                    cell.text = "n. d."
                elif isinstance(value, (int, float)):
                    cell.text = _fmt(value, column.decimals)
                else:
                    cell.text = str(value)[:28]
                para = cell.text_frame.paragraphs[0]
                para.runs[0].font.size = Pt(9)
                if column.type != "text":
                    para.alignment = PP_ALIGN.RIGHT
                if highlight is not None:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = highlight

        if len(table.rows) > MAX_TABLE_ROWS:
            note = slide.shapes.add_textbox(
                Inches(0.6), top + Inches(0.32 * (len(rows) + 1) + 0.1), Inches(12.1), Inches(0.3)
            )
            note.text_frame.text = (
                f"Mostrando {MAX_TABLE_ROWS} de {len(table.rows)} zonas. "
                "La tabla completa está en el informe web y en el Excel."
            )
            note.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
            note.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED

    def _placerank(self, deck: Deck, placerank: dict[str, Any]) -> None:
        slide = self._blank(deck)
        weights = " · ".join(
            f"{k} {int(v * 100)}%" for k, v in (placerank.get("weights") or {}).items()
        )
        self._title(slide, "PlaceRank", f"Índice ponderado — {weights}")

        rows = placerank.get("rows", [])[:MAX_TABLE_ROWS]
        if not rows:
            return
        shape = slide.shapes.add_table(
            len(rows) + 1, 5, Inches(0.6), Inches(1.6), Inches(12.1),
            Inches(0.34 * (len(rows) + 1)),
        )
        grid = shape.table
        for c, title in enumerate(["#", "Zona", "Score", "Categoría", "Match %"]):
            cell = grid.cell(0, c)
            cell.text = title
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(10)
            cell.text_frame.paragraphs[0].runs[0].font.bold = True
            cell.text_frame.paragraphs[0].runs[0].font.color.rgb = WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT

        for r, row in enumerate(rows, start=1):
            values = [
                str(row.get("rank", r)),
                str(row.get("geo_name", "")),
                _fmt(row.get("score", 0), 1),
                str(row.get("category", "")),
                _fmt((row.get("dimensions") or {}).get("match", 0), 1),
            ]
            for c, value in enumerate(values):
                cell = grid.cell(r, c)
                cell.text = value
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(10)

    def _sources(self, deck: Deck, report: Report) -> None:
        slide = self._blank(deck)
        self._title(slide, "Fuentes y metodología")
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.2))
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = "Fuentes de datos"
        frame.paragraphs[0].runs[0].font.bold = True
        frame.paragraphs[0].runs[0].font.size = Pt(14)

        for source in report.sources:
            para = frame.add_paragraph()
            para.text = f"· {source.get('name')} — {source.get('attribution')}"
            para.runs[0].font.size = Pt(11)
            para.runs[0].font.color.rgb = MUTED

        if report.warnings:
            para = frame.add_paragraph()
            para.text = "Avisos"
            para.runs[0].font.bold = True
            para.runs[0].font.size = Pt(14)
            for warning in report.warnings:
                item = frame.add_paragraph()
                item.text = f"· {warning}"
                item.runs[0].font.size = Pt(10)
                item.runs[0].font.color.rgb = MUTED


def _fmt(value: Any, decimals: int = 2) -> str:
    if value is None:
        return "n. d."
    if isinstance(value, (int, float)):
        return f"{value:,.{decimals}f}".replace(",", "␟").replace(".", ",").replace("␟", ".")
    return str(value)
